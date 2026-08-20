"""Generates sample_docs/sample_invoice.{pdf,docx,txt} - three real files with
identical invoice content in different formats, used by tests, the README
examples, and manual API exploration."""
import os

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(HERE, "..", "sample_docs")
os.makedirs(OUT_DIR, exist_ok=True)

VENDOR = "Acme Robotics Inc."
ADDRESS = "123 Innovation Way, Springfield, IL"
INVOICE_NUMBER = "INV-10234"
INVOICE_DATE = "03/14/2026"
DUE_DATE = "04/13/2026"
LINE_ITEMS = [
    ("Widget assembly (Model X)", "10", "45.00", "450.00"),
    ("Custom bracket fabrication", "4", "120.00", "480.00"),
    ("Rush shipping fee", "1", "75.50", "75.50"),
]
TOTAL = "1,005.50"


def build_txt():
    lines = [
        VENDOR,
        ADDRESS,
        "",
        f"Invoice Number: {INVOICE_NUMBER}",
        f"Invoice Date: {INVOICE_DATE}",
        f"Due Date: {DUE_DATE}",
        "",
        "Bill To: Contoso Widgets LLC",
        "",
        f"{'Description':<32}{'Qty':<8}{'Unit Price':<14}{'Amount':<10}",
    ]
    for desc, qty, unit_price, amount in LINE_ITEMS:
        lines.append(f"{desc:<32}{qty:<8}${unit_price:<13}${amount:<10}")
    lines += ["", f"Total: ${TOTAL}", "", "Thank you for your business!"]
    path = os.path.join(OUT_DIR, "sample_invoice.txt")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return path


def build_pdf():
    path = os.path.join(OUT_DIR, "sample_invoice.pdf")
    c = canvas.Canvas(path, pagesize=letter)
    width, height = letter
    y = height - 72
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, y, VENDOR)
    y -= 16
    c.setFont("Helvetica", 10)
    c.drawString(72, y, ADDRESS)
    y -= 30
    c.drawString(72, y, f"Invoice Number: {INVOICE_NUMBER}")
    y -= 16
    c.drawString(72, y, f"Invoice Date: {INVOICE_DATE}")
    y -= 16
    c.drawString(72, y, f"Due Date: {DUE_DATE}")
    y -= 30
    c.drawString(72, y, "Bill To: Contoso Widgets LLC")
    y -= 30

    col_x = [72, 300, 380, 470]
    c.setFont("Helvetica-Bold", 10)
    for x, text in zip(col_x, ["Description", "Qty", "Unit Price", "Amount"]):
        c.drawString(x, y, text)
    y -= 16
    c.setFont("Helvetica", 10)
    for desc, qty, unit_price, amount in LINE_ITEMS:
        for x, text in zip(col_x, [desc, qty, f"${unit_price}", f"${amount}"]):
            c.drawString(x, y, text)
        y -= 16

    y -= 20
    c.setFont("Helvetica-Bold", 11)
    c.drawString(72, y, f"Total: ${TOTAL}")
    y -= 30
    c.setFont("Helvetica", 9)
    c.drawString(72, y, "Thank you for your business!")
    c.save()
    return path


def build_docx():
    path = os.path.join(OUT_DIR, "sample_invoice.docx")
    doc = Document()
    doc.add_paragraph(VENDOR)
    doc.add_paragraph(ADDRESS)
    doc.add_paragraph(f"Invoice Number: {INVOICE_NUMBER}")
    doc.add_paragraph(f"Invoice Date: {INVOICE_DATE}")
    doc.add_paragraph(f"Due Date: {DUE_DATE}")
    doc.add_paragraph("Bill To: Contoso Widgets LLC")

    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Description", "Qty", "Unit Price", "Amount"
    for desc, qty, unit_price, amount in LINE_ITEMS:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = desc, qty, f"${unit_price}", f"${amount}"

    doc.add_paragraph(f"Total: ${TOTAL}")
    doc.add_paragraph("Thank you for your business!")
    doc.save(path)
    return path


if __name__ == "__main__":
    for builder in (build_txt, build_pdf, build_docx):
        path = builder()
        print(f"wrote {path}")
