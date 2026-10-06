"""Text extraction from messy documents: PDF (text layer or scanned), DOCX,
plain text, and images (PNG/JPEG/WebP/TIFF/BMP) of an invoice.

Everything downstream speaks text, so the job of this module is to always
produce text - by reading the PDF text layer, by OCR'ing pixels locally (see
`ocr.py`), or, when OCR is unavailable/unreliable, by handing the page images
to the caller so the AI vision backend can read the pixels instead. That last
case is why `ParsedDocument` carries `images` alongside `text`.
"""
import os
import tempfile
import uuid
from dataclasses import dataclass, field
from typing import List, Optional

import pdfplumber
from docx import Document as DocxDocument

from .ocr import OcrUnavailable, get_engine

TEXT_EXTENSIONS = (".txt", ".text")
IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff")
SUPPORTED_EXTENSIONS = (".pdf", ".docx") + TEXT_EXTENSIONS + IMAGE_EXTENSIONS

# A PDF with fewer than this many characters of text layer is treated as a scan
# (photocopier output, phone photo printed to PDF) and goes through OCR.
PDF_TEXT_MIN_CHARS = 40

# Page render size for OCR/vision, as a PDF points multiplier (72 dpi * scale).
# ~200 dpi is the sweet spot for RapidOCR: below it small table digits blur.
DEFAULT_RENDER_SCALE = 2.8

# Guard against a 500-page scan eating a serverless container's memory.
DEFAULT_MAX_PAGES = 8


class UnsupportedDocumentType(ValueError):
    pass


@dataclass
class ParsedDocument:
    text: str
    # "pdf-text" | "docx" | "txt" | "ocr" | "image" (no usable text, vision needed)
    source: str
    ocr_engine: Optional[str] = None
    ocr_score: Optional[float] = None
    # image files (the upload itself, or rendered PDF pages) for the vision path
    images: List[str] = field(default_factory=list)
    notes: List[str] = field(default_factory=list)

    @property
    def is_image(self) -> bool:
        return bool(self.images)


def extension_of(file_path: str) -> str:
    return os.path.splitext(file_path)[1].lower()


def extract_text(file_path: str) -> str:
    """Back-compatible entry point: just the text. Use `parse_document` when
    the OCR engine, page images or provenance matter."""
    return parse_document(file_path).text


def parse_document(file_path: str) -> ParsedDocument:
    ext = extension_of(file_path)
    if ext == ".pdf":
        return _parse_pdf(file_path)
    if ext == ".docx":
        return ParsedDocument(text=_extract_docx(file_path), source="docx")
    if ext in TEXT_EXTENSIONS:
        return ParsedDocument(text=_extract_txt(file_path), source="txt")
    if ext in IMAGE_EXTENSIONS:
        return _from_images([file_path])
    raise UnsupportedDocumentType(f"unsupported document type: {ext or '(no extension)'}")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _extract_pdf(file_path: str) -> str:
    pages = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            # layout=True preserves horizontal whitespace between columns,
            # which the heuristic line-item parser relies on.
            pages.append(page.extract_text(layout=True) or "")
    return "\n".join(pages).strip()


def _parse_pdf(file_path: str) -> ParsedDocument:
    text = _extract_pdf(file_path)
    if len(text) >= PDF_TEXT_MIN_CHARS:
        return ParsedDocument(text=text, source="pdf-text")

    # No (or near-empty) text layer: it is a scan. Render and OCR the pages.
    images = render_pdf_pages(file_path)
    parsed = _from_images(images)
    parsed.notes.insert(0, f"pdf has no text layer ({len(text)} chars); OCR'd {len(images)} page(s)")
    if not parsed.text:
        parsed.notes.insert(0, "no text layer and no OCR text - needs the AI vision backend")
    return parsed


def render_pdf_pages(file_path: str, scale: Optional[float] = None, max_pages: Optional[int] = None) -> List[str]:
    """Rasterize each PDF page to a PNG and return the image paths. Pages are
    written to the system temp dir (override with INVOICELENS_PAGE_DIR) rather
    than next to the source file, so committed samples and read-only mounts
    stay clean and Vercel's /tmp-only rule is satisfied. Uses pypdfium2 (a
    pdfplumber dependency), which ships its own PDF renderer - no system binary
    needed."""
    import pypdfium2 as pdfium

    if scale is None:
        scale = float(os.environ.get("OCR_RENDER_SCALE", DEFAULT_RENDER_SCALE))
    if max_pages is None:
        max_pages = int(os.environ.get("OCR_MAX_PAGES", DEFAULT_MAX_PAGES))

    out_dir = os.environ.get("INVOICELENS_PAGE_DIR") or tempfile.gettempdir()
    os.makedirs(out_dir, exist_ok=True)
    # Per-document uid in the filename: two different "scan.pdf" files rendered
    # into the same shared temp dir would otherwise overwrite each other.
    stem = f"{os.path.splitext(os.path.basename(file_path))[0]}-{uuid.uuid4().hex[:8]}"
    paths: List[str] = []
    document = pdfium.PdfDocument(file_path)
    try:
        for index in range(min(len(document), max_pages)):
            image = document[index].render(scale=scale).to_pil()
            out_path = os.path.join(out_dir, f"{stem}_page{index + 1}.png")
            image.save(out_path)
            paths.append(out_path)
    finally:
        document.close()
    return paths


# ---------------------------------------------------------------------------
# Images / OCR
# ---------------------------------------------------------------------------

def _from_images(images: List[str]) -> ParsedDocument:
    """OCR the given page images when an engine is installed. The images are
    always kept on the result so the caller can escalate to a vision model."""
    if not images:
        return ParsedDocument(text="", source="image", notes=["no pages to read"])
    try:
        engine = get_engine()
    except OcrUnavailable as exc:
        return ParsedDocument(text="", source="image", images=images, notes=[str(exc)])

    result = engine.run(images)
    notes = [f"ocr engine {result.engine} on {len(images)} page(s), mean score {result.mean_score}"]
    if not result.text:
        notes.append("OCR found no text")
    elif result.unreliable:
        notes.append(f"OCR mean score {result.mean_score} below {engine.name} reliability floor")
    return ParsedDocument(
        text=result.text,
        source="ocr",
        ocr_engine=result.engine,
        ocr_score=result.mean_score,
        images=images,
        notes=notes,
    )


# ---------------------------------------------------------------------------
# DOCX / TXT
# ---------------------------------------------------------------------------

def _extract_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()


def _extract_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read().strip()
